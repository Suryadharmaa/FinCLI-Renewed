"""V3 behavioral regressions: finance invariants, boundaries, concurrency and shared contracts."""

from __future__ import annotations

import asyncio
import base64
import io
import json
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from fincli.app.cli.router import CommandRouter
from fincli.app.providers.ai.base import AIResponse
from fincli.app.providers.market.base import Candle
from fincli.app.storage.config import ConfigManager
from fincli.app.storage.database import FinCLIDatabase
from fincli.app.web.bridge import execute_command
from fincli.app.workspace.company import CompanyService, dcf, ratios
from fincli.app.workspace.jobs import JobManager
from fincli.app.workspace.models import Instrument, finite, json_safe
from fincli.app.workspace.portfolio import compare_fixed_holdings, portfolio_intelligence
from fincli.app.workspace.screener import ScreenExpression
from fincli.app.workspace.store import WorkspaceStore
from fincli.app.workspace.tradingview import TradingViewBridge
from tests.test_command_smoke import SmokeAIProvider, SmokeMarketProvider
from tests.workspace_support import company_data


@pytest.fixture
def router(tmp_path):
    instance = CommandRouter(
        config=ConfigManager(tmp_path / "config.json"),
        db=FinCLIDatabase(tmp_path / "finance.db"),
        market_provider=SmokeMarketProvider(),
        ai_provider=SmokeAIProvider(),
    )
    instance.workspace_service.company.loader = company_data
    yield instance
    instance.shutdown()


def test_company_ratios_align_dates_and_withhold_missing_inputs():
    data = company_data("AAPL")
    data["statements"]["income"].reverse()
    rows = ratios(data["statements"])
    assert rows[0]["period"] == "2025-12-31"
    assert rows[0]["revenue_growth"] == pytest.approx(20)
    assert rows[0]["net_margin"] == pytest.approx(100 / 6)
    assert rows[0]["roe"] == pytest.approx(20)
    assert rows[0]["roic"] == pytest.approx(24 / 130 * 100)
    assert rows[-1]["roe"] is None
    assert rows[-1]["roic"] is None
    data["statements"]["income"][1]["values"].pop("Tax Provision")
    assert ratios(data["statements"])[0]["roic"] is None


def test_peer_alignment_withholds_incompatible_fiscal_years():
    def loader(symbol):
        data = company_data(symbol)
        if symbol == "MSFT":
            data["ratios"] = [{**r, "period": r["period"].replace("12-31", "06-30")} for r in data["ratios"]]
        return data

    output = asyncio.run(CompanyService(loader=loader).peers(["AAPL", "MSFT"]))
    assert output["status"] == "partial"
    assert output["data"]["period"] is None
    assert all(r["metrics"] is None for r in output["data"]["rows"])


def test_company_provider_failure_is_partial():
    def unavailable(symbol):
        raise ConnectionError("provider failed")

    output = asyncio.run(CompanyService(loader=unavailable).get("NASDAQ:AAPL"))
    assert output["status"] == "partial"
    assert output["symbol"] == "NASDAQ:AAPL"
    assert "ConnectionError" in output["warnings"][0]


def test_dcf_constant_cashflow_and_sensitivity():
    data = dcf(100, 0, 0.10, 0, cash=50, debt=20, shares=10)["data"]
    assert data["enterprise_value"] == pytest.approx(1000)
    assert data["equity_value"] == pytest.approx(1030)
    assert data["per_share"] == pytest.approx(103)
    assert len(data["sensitivity"]) == 9
    by_discount = sorted((r for r in data["sensitivity"] if r["growth"] == 0), key=lambda r: r["discount"])
    assert by_discount[0]["per_share"] > by_discount[-1]["per_share"]


@pytest.mark.parametrize(
    "params",
    [
        dict(fcf=-1),
        dict(discount=0.02),
        dict(shares=0),
        dict(fcf=float("inf")),
        dict(growth=float("nan")),
        dict(years=2.5),
        dict(years=True),
        dict(cash=-1),
        dict(fcf=1e308),
    ],
)
def test_dcf_rejects_invalid_or_overflowing_inputs(params):
    inputs = dict(fcf=100, growth=0.05, discount=0.1, terminal=0.02)
    inputs.update(params)
    with pytest.raises(ValueError):
        dcf(**inputs)


def test_currency_attribution_and_stress_conserve_pnl():
    positions = [{"symbol": "EURCO", "currency": "EUR", "quantity": 2, "average_price": 10}]
    data = portfolio_intelligence(
        positions, {"EURCO": 12}, fx={"EUR": 1.2}, entry_fx={"EUR": 1.1}, shock=-0.2, fx_shock=-0.1
    )["data"]
    row = data["positions"][0]
    assert row["price_contribution"] + row["fx_contribution"] == pytest.approx(row["base_unrealized_pnl"])
    assert data["total_value"] == pytest.approx(28.8)
    assert data["scenario"]["stressed_valued_total"] == pytest.approx(28.8 * 0.8 * 0.9)
    incomplete = portfolio_intelligence(positions, {"EURCO": 12})["data"]
    assert incomplete["total_value"] is None and incomplete["rebalance_preview"] == []
    no_entry = portfolio_intelligence(positions, {"EURCO": 12}, fx={"EUR": 1.2})["data"]
    assert no_entry["unrealized_pnl"] is None


def test_rebalance_preview_is_cash_neutral_before_fees():
    positions = [{"symbol": s, "quantity": q, "average_price": 10, "currency": "USD"} for s, q in (("A", 1), ("B", 3))]
    data = portfolio_intelligence(positions, {"A": 10, "B": 10}, fee_bps=10)["data"]
    assert sum(r["delta_base"] for r in data["rebalance_preview"]) == pytest.approx(0)
    assert data["estimated_fees"] == pytest.approx(0.02)


def test_benchmark_uses_same_dates_and_currency():
    positions = [{"symbol": "A", "quantity": 2, "currency": "USD"}]
    histories = {
        "A": {"2025-01-01": 1, "2025-01-02": 10, "2025-01-03": 12},
        "SPY": {"2025-01-02": 100, "2025-01-03": 110, "2025-01-04": 500},
    }
    data = compare_fixed_holdings(positions, histories, "SPY", "USD")
    assert data["start"] == "2025-01-02" and data["end"] == "2025-01-03"
    assert data["portfolio_return_pct"] == pytest.approx(20)
    assert data["benchmark_return_pct"] == pytest.approx(10)
    assert compare_fixed_holdings(positions, histories, "SPY", "EUR")["status"] == "partial"


@pytest.mark.parametrize(
    "expression",
    [
        '__import__("os").system("x")',
        "price > 0; exit()",
        "price > nan",
        "unknown > 2",
        "pe > 2%",
        "(price > 0)",
        "price > 1 and ",
    ],
)
def test_screener_rejects_code_and_ambiguous_units(expression):
    with pytest.raises(ValueError):
        ScreenExpression(expression)


def test_screener_precedence_and_missing_fields_are_explicit():
    expression = ScreenExpression("revenue_growth > 10% and rsi < 40 or net_margin > 20%")
    assert expression.evaluate({"revenue_growth": 20, "rsi": 30})["passed"]
    missing = expression.evaluate({"rsi": 30})
    assert not missing["passed"] and "revenue_growth" in missing["missing_fields"]
    assert not ScreenExpression("price != 0").evaluate({"price": None})["passed"]


def test_document_evidence_is_traceable_to_original_page(router):
    docs = router.workspace_service.documents
    raw = b"Revenue grew.\fOperating margin risk from higher input costs."
    imported = docs.import_bytes("../report.txt", raw, "AAPL")["data"]
    again = docs.import_bytes("report.txt", raw, "AAPL")["data"]
    assert imported["id"] == again["id"]
    hit = docs.search("margin risk", "AAPL")["data"]["evidence"][0]
    page = docs.page(hit["document_id"], hit["page"])
    assert hit["page"] == 2 and hit["quote"] in page["text"]
    assert hit["digest"] == imported["digest"]
    assert docs.search("margin risk", "MSFT")["data"]["evidence"] == []
    with pytest.raises(ValueError):
        docs.page(imported["id"], 3)


def test_pdf_text_extraction_and_invalid_files(router):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 10 100 Td (Margin risk evidence) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    buffer = io.BytesIO()
    writer.write(buffer)
    docs = router.workspace_service.documents
    doc = docs.import_bytes("report.pdf", buffer.getvalue())["data"]
    assert "Margin risk" in docs.page(doc["id"], 1)["text"]
    for title, raw in (("bad.pdf", b"invalid PDF"), ("bad.txt", b"\xff"), ("image.txt", b" "), ("bad.exe", b"code")):
        with pytest.raises(ValueError):
            docs.import_bytes(title, raw)


@pytest.mark.parametrize(
    "payload", [{"panels": []}, {"panels": ["unsupported"]}, {"symbols": ["<script>"]}, {"interval": "invalid"}]
)
def test_layout_validation_preserves_previous_layout(router, payload):
    store = router.workspace_service.store
    original = store.save_layout("research", {"symbol": "AAPL", "panels": ["chart", "company"]})
    with pytest.raises(ValueError):
        store.save_layout("research", payload)
    assert store.layouts()[0] == original


def test_layout_and_runs_survive_service_restart(router):
    store = router.workspace_service.store
    store.save_layout("custom", {"symbols": ["AAPL", "MSFT"], "symbol": "MSFT", "panels": ["news"]})
    screen = router.workspace_service.screen(["AAPL"], "revenue_growth > 10%")
    reopened = WorkspaceStore(router.db)
    assert reopened.layouts()[-1]["symbol"] == "MSFT"
    assert reopened.get_run(screen["data"]["id"]) == screen


def wait_job(manager, job_id):
    until = time.monotonic() + 3
    while time.monotonic() < until:
        state = manager.get(job_id)
        if state["status"] in {"completed", "failed", "cancelled"}:
            return state
        threading.Event().wait(0.01)
    pytest.fail("Job did not complete")


def test_job_cancellation_and_capacity_do_not_publish_results():
    manager = JobManager(workers=1, capacity=2)
    started, release = threading.Event(), threading.Event()

    def operation(context):
        started.set()
        release.wait(2)
        context.checkpoint(50, "provider done")
        return {"should_not_publish": True}

    try:
        first = manager.submit(operation, "running")
        assert started.wait(2)
        second = manager.submit(lambda ctx: {"queued": True}, "queued")
        with pytest.raises(ValueError):
            manager.submit(operation, "full")
        manager.cancel(first["id"])
        manager.cancel(second["id"])
        release.set()
        for job in (first, second):
            state = wait_job(manager, job["id"])
            assert state["status"] == "cancelled" and state["result"] is None
    finally:
        release.set()
        manager.shutdown()


def test_shared_event_loop_handles_concurrent_job_threads(router):
    async def value(i):
        await asyncio.sleep(0.01)
        return i

    with ThreadPoolExecutor(max_workers=6) as pool:
        result = list(pool.map(lambda i: router.market_service.run(value(i)), range(12)))
    assert result == list(range(12))
    with ThreadPoolExecutor(max_workers=6) as pool:
        services = list(pool.map(lambda _: router.workspace_service, range(12)))
    assert all(service is services[0] for service in services)


@pytest.mark.parametrize(
    "plan",
    [
        {"steps": [{"tool": "company", "params": {"symbol": "AAPL"}}, {"tool": "order", "params": {}}]},
        {"steps": [{"tool": "company", "params": {"symbol": "AAPL", "path": "/tmp/file"}}]},
        {"steps": [{"tool": "market", "params": {}}]},
        {"steps": [{"tool": "peers", "params": {"symbols": "AAPL,MSFT"}}]},
    ],
)
def test_ai_plan_is_fully_validated_before_tools(router, monkeypatch, plan):
    async def complete(request):
        return AIResponse("fixture", "fixture", json.dumps(plan))

    monkeypatch.setattr(router.ai_provider, "complete", complete)
    monkeypatch.setattr(router.workspace_service, "request", lambda *args: pytest.fail("Invalid plan executed a tool"))
    with pytest.raises(ValueError):
        router.workspace_service.workflow("Research AAPL")


def test_workflow_keeps_provenance_and_document_citations(router, monkeypatch):
    router.workspace_service.documents.import_bytes("risk.md", b"Margin risk from input costs.", "AAPL")
    prompts = []

    async def complete(request):
        prompts.append(request.prompt)
        return AIResponse(
            "fixture",
            "fixture",
            json.dumps({"steps": [{"tool": "documents", "params": {"query": "margin risk", "symbol": "AAPL"}}]})
            if len(prompts) == 1
            else "Margin risk observed [1:D1].",
        )

    monkeypatch.setattr(router.ai_provider, "complete", complete)
    output = router.workspace_service.workflow("Explain AAPL margin risk")
    assert output["data"]["read_only"]
    assert output["data"]["steps"][0]["result"]["data"]["evidence"][0]["citation_id"] == "D1"
    assert "untrusted evidence" in prompts[-1] and '"provenance"' in prompts[-1]
    assert router.workspace_service.store.get_run(output["data"]["id"]) == output


def test_cli_web_use_identical_structured_company_contract(router):
    cli = router.route("/company AAPL")
    web = execute_command(router, "/company AAPL")
    assert cli.status == "ready" and web.ok
    assert cli.metadata["workspace"]["data"] == web.metadata["workspace"]["data"]
    assert web.metadata["workspace"]["contract"] == "3.0"
    assert not execute_command(router, "/document import /tmp/private.txt").ok
    assert not execute_command(router, "/workspace export id /tmp/private.json").ok


def test_screen_excel_contains_results_and_provenance(router):
    output = router.workspace_service.screen(["AAPL", "MSFT"], "revenue_growth > 10%")
    book = load_workbook(io.BytesIO(router.workspace_service.excel_bytes(output["data"]["id"])))
    assert book["Screen"].max_row == 3
    assert book["Screen"]["A2"].value == "AAPL"
    assert "Provenance" in book.sheetnames


def test_experiment_saves_full_input_snapshot(router, monkeypatch):
    candles = [
        Candle(datetime(2025, 1, 1, tzinfo=UTC) + timedelta(days=i), 100 + i, 101 + i, 99 + i, 100 + i, 1000)
        for i in range(80)
    ]

    async def history(*args, **kwargs):
        return candles

    monkeypatch.setattr(router.market_service, "history", history)
    output = router.workspace_service.backtest({"symbol": "AAPL"})
    saved = router.workspace_service.store.get_run(output["data"]["id"])
    assert len(saved["data"]["input_candles"]) == 80
    assert saved["data"]["parameters"]["strategy"] == "sma_cross"


def test_tradingview_allowlist_confirmation_and_no_shell(tmp_path, monkeypatch):
    script = tmp_path / "index.js"
    script.write_text("// external fixture", encoding="utf-8")
    bridge = TradingViewBridge(str(script))
    assert bridge.capabilities()["data"]["orders"] is False
    with pytest.raises(ValueError):
        bridge.execute("symbol", {"symbol": "AAPL"})
    for action in ("eval", "extract", "order"):
        with pytest.raises(ValueError):
            bridge.execute(action, confirmed=True)
    calls = []

    def run(args, **kwargs):
        calls.append((args, kwargs))
        return type("Completed", (), {"returncode": 0, "stdout": '{"success":true}'})()

    monkeypatch.setattr("fincli.app.workspace.tradingview.subprocess.run", run)
    bridge.execute("draw", {"price": 100}, confirmed=True)
    args, kwargs = calls[-1]
    assert args[2:6] == ["draw", "shape", "--type", "horizontal_line"]
    assert kwargs["shell"] is False and kwargs["env"]["TV_CDP_HOST"] == "127.0.0.1"


def test_authenticated_api_jobs_documents_and_exports(router, monkeypatch):
    import fincli.app.web.api as api_module

    monkeypatch.setenv("FINCLI_DESKTOP_TOKEN", "test-v3-local-token")
    monkeypatch.setattr(api_module, "ConfigManager", lambda: router.config)
    monkeypatch.setattr(api_module, "FinCLIDatabase", lambda: router.db)
    monkeypatch.setattr(api_module, "CommandRouter", lambda **kwargs: router)
    headers = {"Authorization": "Bearer test-v3-local-token", "X-FinCLI-CSRF": "local-web"}
    with TestClient(api_module.create_app()) as client:
        assert client.get("/api/workspace/layouts").status_code == 401
        assert (
            client.post("/api/workspace/jobs", headers={"Authorization": headers["Authorization"]}, json={}).status_code
            == 403
        )
        assert client.post("/api/workspace/jobs", headers=headers, json={"action": "order"}).status_code == 422
        assert (
            client.post(
                "/api/workspace/layouts",
                headers=headers,
                json={"name": "custom", "panels": ["chart"], "symbol": "AAPL"},
            ).status_code
            == 200
        )
        uploaded = client.post(
            "/api/workspace/documents",
            headers=headers,
            json={
                "title": "report.txt",
                "content": base64.b64encode(b"Margin risk evidence").decode(),
                "symbol": "AAPL",
            },
        ).json()
        assert (
            client.get(f"/api/workspace/documents/{uploaded['data']['id']}/1", headers=headers).json()["text"]
            == "Margin risk evidence"
        )
        assert (
            client.post(
                "/api/workspace/documents", headers=headers, json={"title": "bad.txt", "content": "###"}
            ).status_code
            == 422
        )
        job = client.post(
            "/api/workspace/jobs",
            headers=headers,
            json={"action": "screen", "params": {"symbols": ["AAPL"], "expression": "revenue_growth > 10%"}},
        ).json()
        state = wait_job(router.workspace_service.jobs, job["id"])
        assert state["status"] == "completed"
        assert (
            client.get(f"/api/workspace/job-status?ids={job['id']}", headers=headers).json()["jobs"][job["id"]][
                "status"
            ]
            == "completed"
        )
        run_id = state["result"]["data"]["id"]
        export = client.get(f"/api/workspace/runs/{run_id}/export", headers=headers)
        assert export.status_code == 200 and export.content.startswith(b"PK")
        assert client.get("/api/workspace/runs/missing/export", headers=headers).status_code == 422
        assert (
            client.post("/api/workspace/tradingview", headers=headers, json={"action": "capabilities"}).json()["data"][
                "orders"
            ]
            is False
        )


def test_identity_and_json_boundary():
    instrument = Instrument.parse("NASDAQ:AAPL")
    assert instrument.exchange == "NASDAQ" and instrument.provider_symbol == "AAPL"
    assert finite(True) is None and finite("nan") is None
    assert json_safe({"a": float("inf")}) == {"a": None}


def test_new_commands_preserve_aliases(router):
    assert router.route("/p intelligence").metadata["workspace"]["kind"] == "portfolio_intelligence"
    assert router.route('/s AAPL --where "revenue_growth > 10%"').metadata["workspace"]["kind"] == "screen"


def test_per_share_overflow_is_rejected():
    with pytest.raises(ValueError):
        dcf(100, 0.05, 0.1, 0.02, shares=1e-308)


def test_web_confirmation_reaches_tradingview_adapter(router, monkeypatch):
    from fincli.app.workspace.models import result

    calls = []

    def execute(action, params, confirmed=False):
        calls.append(confirmed)
        return result("tradingview", {"success": True})

    monkeypatch.setattr(router.workspace_service.tradingview, "execute", execute)
    blocked = execute_command(router, "/tv symbol AAPL")
    assert blocked.status == "confirmation_required" and not calls
    accepted = execute_command(router, "/tv symbol AAPL", confirmed=True)
    assert accepted.ok and calls == [True]


def test_ai_validates_all_enum_arguments_before_providers(router, monkeypatch):
    async def complete(request):
        return AIResponse(
            "fixture",
            "fixture",
            json.dumps(
                {
                    "steps": [
                        {"tool": "company", "params": {"symbol": "AAPL"}},
                        {"tool": "backtest", "params": {"symbol": "MSFT", "strategy": "arbitrary"}},
                    ]
                }
            ),
        )

    monkeypatch.setattr(router.ai_provider, "complete", complete)
    monkeypatch.setattr(
        router.workspace_service, "request", lambda *args: pytest.fail("Invalid plan reached a provider")
    )
    with pytest.raises(ValueError):
        router.workspace_service.workflow("Research AAPL then test MSFT")


def test_saved_workflow_evidence_and_coverage_are_visible_in_cli(router):
    from rich.console import Console

    workflow = router.workspace_service.workflow("Research AAPL")
    command = router.route(f"/workspace run {workflow['data']['id']}")
    console = Console(record=True, width=200)
    console.print(command.renderable)
    text = console.export_text()
    assert "Total Revenue" in text and '"provenance"' in text
    valuation = router.route("/valuation 100 0.05 0.1 0.02")
    console.print(valuation.renderable)
    assert "User-supplied unlevered FCF" in console.export_text()
